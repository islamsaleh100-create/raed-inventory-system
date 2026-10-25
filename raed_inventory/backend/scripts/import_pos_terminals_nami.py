#!/usr/bin/env python3
"""Import Nami POS terminals from Import_Preview sheet (local/test only)."""
from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path
from typing import Any

import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from app.models import Branch, User  # noqa: E402
from app.models.pos_terminal import PosTerminal  # noqa: E402
from app.schemas.pos_terminal import PosTerminalCreate  # noqa: E402
from app.services import pos_terminal_service  # noqa: E402
from app.services.pos_terminal_utils import normalize_terminal_text  # noqa: E402

COL_BRANCH_LABEL = 3  # D
COL_BRANCH_FINAL = 6  # G
COL_TERMINAL = 7  # H
COL_PROVIDER = 8  # I
COL_STATUS = 9  # J
COL_NOTES = 10  # K
SHEET = "Import_Preview"
AUDITOR_USERNAME = "audit.officer"
_DASH_LABELS = frozenset({"—", "-", "–", "−"})


def _sheet_branch_label(raw: Any) -> str | None:
    if raw is None:
        return None
    text = str(raw).strip()
    if not text or text in _DASH_LABELS:
        return None
    return text


def _assert_local_db(url: str) -> None:
    lower = url.lower()
    blocked = ("rlwy.net", "railway.app", "railway.internal", "amazonaws.com", "proxy.rlwy")
    if any(b in lower for b in blocked):
        raise SystemExit("Refusing to run: DATABASE_URL looks like a non-local host.")
    if "raed_inventory_test" not in lower and "localhost" not in lower and "127.0.0.1" not in lower:
        raise SystemExit("Refusing to run: DATABASE_URL must be local test (localhost / raed_inventory_test).")


def _assert_db_guard(url: str, *, allow_remote: bool) -> None:
    if allow_remote:
        if os.environ.get("POS_IMPORT_ALLOW_REMOTE") != "YES":
            raise SystemExit("Refusing remote: set POS_IMPORT_ALLOW_REMOTE=YES")
        typed = sys.stdin.readline().strip()
        if typed != "IMPORT":
            raise SystemExit("Refusing remote: confirmation word must be IMPORT")
        return
    _assert_local_db(url)


def _branch_map(db) -> dict[str, int]:
    rows = db.query(Branch.id, Branch.branch_name).filter(Branch.is_deleted == False).all()  # noqa: E712
    out: dict[str, int] = {}
    for bid, name in rows:
        key = (name or "").strip().casefold()
        if key:
            out[key] = bid
    return out


def _parse_rows(path: Path) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"Sheet {SHEET!r} not found in {path}")
    ws = wb[SHEET]
    parsed: list[dict[str, Any]] = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        terminal_raw = row[COL_TERMINAL] if len(row) > COL_TERMINAL else None
        terminal = normalize_terminal_text(str(terminal_raw or ""))
        provider = str(row[COL_PROVIDER] or "").strip().lower()
        status = str(row[COL_STATUS] or "").strip().lower()
        branch_final = str(row[COL_BRANCH_FINAL] or "").strip()
        branch_label = _sheet_branch_label(row[COL_BRANCH_LABEL] if len(row) > COL_BRANCH_LABEL else None)
        notes = str(row[COL_NOTES] or "").strip() or None
        parsed.append(
            {
                "row": idx,
                "terminal_number": terminal,
                "provider": provider,
                "status": status,
                "branch_final": branch_final,
                "branch_label": branch_label,
                "notes": notes,
            }
        )
    return parsed


def _load_existing_keys(db) -> set[tuple[str, str]]:
    rows = db.query(PosTerminal.provider, PosTerminal.terminal_number).all()
    return {(p, n) for p, n in rows}


def main() -> None:
    parser = argparse.ArgumentParser(description="Import Nami POS terminals from Excel")
    parser.add_argument("workbook", type=Path, help="Path to POS_Nami_import_v2.xlsx")
    parser.add_argument("--apply", action="store_true", help="Write to database (default: dry-run)")
    parser.add_argument(
        "--allow-remote",
        action="store_true",
        help="Allow non-local DATABASE_URL when POS_IMPORT_ALLOW_REMOTE=YES and stdin confirms IMPORT",
    )
    args = parser.parse_args()

    url = os.environ.get("DATABASE_URL", "").strip()
    if not url:
        raise SystemExit("DATABASE_URL is required")
    _assert_db_guard(url, allow_remote=args.allow_remote)

    if not args.workbook.is_file():
        raise SystemExit(f"Workbook not found: {args.workbook}")

    rows = _parse_rows(args.workbook)
    engine = create_engine(url)
    Session = sessionmaker(bind=engine)
    db = Session()
    try:
        auditor = db.query(User).filter(User.username == AUDITOR_USERNAME, User.is_deleted == False).first()  # noqa: E712
        if not auditor:
            raise SystemExit(f"User {AUDITOR_USERNAME!r} not found in database")

        branches = _branch_map(db)
        existing = _load_existing_keys(db)
        seen: set[tuple[str, str]] = set()
        to_create: list[PosTerminalCreate] = []
        rejected: list[str] = []
        with_branch = 0
        without_branch = 0

        for item in rows:
            reasons: list[str] = []
            if not item["terminal_number"]:
                reasons.append("empty terminal number")
            if item["provider"] not in ("hala", "nami", "riyad_bank"):
                reasons.append(f"invalid provider {item['provider']!r}")
            if item["status"] not in ("active", "disabled", "withdrawn"):
                reasons.append(f"invalid status {item['status']!r}")
            key = (item["provider"], item["terminal_number"])
            if item["terminal_number"]:
                if key in seen:
                    reasons.append("duplicate in file")
                if key in existing:
                    reasons.append("duplicate in database")
            branch_id = None
            if item["branch_final"]:
                bid = branches.get(item["branch_final"].casefold())
                if bid is None:
                    reasons.append(f"unknown branch {item['branch_final']!r}")
                else:
                    branch_id = bid
            if reasons:
                rejected.append(f"row {item['row']}: {', '.join(reasons)}")
                continue
            if branch_id is not None:
                with_branch += 1
            else:
                without_branch += 1
            seen.add(key)
            to_create.append(
                PosTerminalCreate(
                    provider=item["provider"],
                    terminal_number=item["terminal_number"],
                    branch_id=branch_id,
                    branch_label=item["branch_label"],
                    status=item["status"],
                    notes=item["notes"],
                )
            )

        print(f"rows_in_sheet={len(rows)}")
        print(f"would_create={len(to_create)}")
        print(f"rejected={len(rejected)}")
        print(f"with_branch={with_branch}")
        print(f"without_branch={without_branch}")
        for line in rejected[:20]:
            print(f"  reject: {line}")
        if len(rejected) > 20:
            print(f"  ... and {len(rejected) - 20} more")

        if rejected:
            raise SystemExit(2)

        if not args.apply:
            print("dry_run=OK (no writes)")
            return

        try:
            for payload in to_create:
                pos_terminal_service.create_terminal(
                    db,
                    user_id=auditor.id,
                    payload=payload,
                    ip_address=None,
                    commit=False,
                )
            db.commit()
        except Exception:
            db.rollback()
            raise
        print(f"apply=OK created={len(to_create)}")
    finally:
        db.close()
        engine.dispose()


if __name__ == "__main__":
    main()
